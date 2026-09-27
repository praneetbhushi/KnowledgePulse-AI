import fitz
from pypdf import PdfReader
from docx import Document as DocxDocument
from pptx import Presentation
from openpyxl import load_workbook


class DocumentProcessor:

    def extract_text(
        self,
        file_path: str,
        file_type: str
    ) -> str:
        """
        Extract text based on document type.
        """

        file_type = file_type.lower()

        if file_type == "pdf":
            return self._extract_pdf(file_path)

        elif file_type == "docx":
            return self._extract_docx(file_path)

        elif file_type == "txt":
            return self._extract_txt(file_path)

        elif file_type == "pptx":
            return self._extract_pptx(file_path)

        elif file_type == "xlsx":
            return self._extract_xlsx(file_path)

        else:
            raise ValueError(
                f"Unsupported file type: {file_type}"
            )

    def _extract_pdf(
        self,
        file_path: str
    ) -> str:

        text = ""

        with fitz.open(file_path) as document:

            for page in document:
                text += page.get_text()
                text += "\n"

        return text.strip()

    def _extract_docx(
        self,
        file_path: str
    ) -> str:

        document = DocxDocument(file_path)

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        return "\n".join(
            paragraphs
        )

    def _extract_txt(
        self,
        file_path: str
    ) -> str:

        with open(
            file_path,
            "r",
            encoding="utf-8"
        ) as file:

            return file.read().strip()

    def _extract_pptx(
        self,
        file_path: str
    ) -> str:

        presentation = Presentation(
            file_path
        )

        text_parts = []

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text = shape.text.strip()

                    if text:
                        text_parts.append(text)

        return "\n".join(
            text_parts
        )

    def _extract_xlsx(
        self,
        file_path: str
    ) -> str:

        workbook = load_workbook(
            file_path,
            data_only=True
        )

        text_parts = []

        for sheet in workbook.worksheets:

            text_parts.append(
                f"Sheet: {sheet.title}"
            )

            for row in sheet.iter_rows(
                values_only=True
            ):

                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]

                if values:

                    text_parts.append(
                        " | ".join(values)
                    )

        return "\n".join(
            text_parts
        )


document_processor = DocumentProcessor()